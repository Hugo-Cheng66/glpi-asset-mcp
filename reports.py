from __future__ import annotations

import csv
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any


DEFAULT_COLUMNS = [
    "id",
    "name",
    "serial",
    "otherserial",
    "uuid",
    "contact",
    "locations_id",
    "states_id",
    "manufacturers_id",
    "computermodels_id",
    "computertypes_id",
    "comment",
    "date_mod",
]


def generate_report(
    items: list[dict[str, Any]],
    reports_dir: Path,
    *,
    report_type: str,
    file_format: str,
    columns: list[str] | None = None,
) -> dict[str, Any]:
    reports_dir.mkdir(parents=True, exist_ok=True)
    safe_type = "".join(ch for ch in report_type if ch.isalnum() or ch in ("-", "_")).strip("_") or "assets"
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    report_id = f"{safe_type}-{timestamp}-{uuid.uuid4().hex[:8]}"
    selected_columns = columns or infer_columns(items)

    if file_format == "csv":
        path = reports_dir / f"{report_id}.csv"
        write_csv(path, items, selected_columns)
    elif file_format == "xlsx":
        path = reports_dir / f"{report_id}.xlsx"
        write_xlsx(path, items, selected_columns)
    else:
        raise ValueError("file_format must be csv or xlsx")

    return {
        "report_id": report_id,
        "format": file_format,
        "path": str(path),
        "rows": len(items),
        "columns": selected_columns,
    }


def generate_custom_report(
    items: list[dict[str, Any]],
    reports_dir: Path,
    *,
    report_type: str,
    file_format: str,
    fields: list[dict[str, str]],
) -> dict[str, Any]:
    rows = project_fields(items, fields)
    columns = [field.get("label") or field["path"] for field in fields]
    return generate_report(
        rows,
        reports_dir,
        report_type=report_type,
        file_format=file_format,
        columns=columns,
    )


def project_fields(items: list[dict[str, Any]], fields: list[dict[str, str]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for item in items:
        row: dict[str, Any] = {}
        for field in fields:
            path = field["path"]
            label = field.get("label") or path
            row[label] = get_path_value(item, path)
        rows.append(row)
    return rows


def expand_items(items: list[dict[str, Any]], expand_path: str | None) -> list[dict[str, Any]]:
    if not expand_path:
        return items

    expanded: list[dict[str, Any]] = []
    for item in items:
        children = get_path_raw(item, expand_path)
        if isinstance(children, dict):
            children = [children]
        if not isinstance(children, list) or not children:
            expanded.append({"asset": item, "item": {}})
            continue
        for child in children:
            if isinstance(child, dict):
                expanded.append({"asset": item, "item": child})
            else:
                expanded.append({"asset": item, "item": {"value": child}})
    return expanded


def infer_columns(items: list[dict[str, Any]]) -> list[str]:
    seen = set()
    columns: list[str] = []
    for column in DEFAULT_COLUMNS:
        if any(column in item for item in items):
            columns.append(column)
            seen.add(column)
    for item in items:
        for key, value in item.items():
            if key not in seen and isinstance(value, (str, int, float, bool, type(None))):
                columns.append(key)
                seen.add(key)
    return columns or DEFAULT_COLUMNS


def get_path_value(item: dict[str, Any], path: str) -> str:
    return _stringify(get_path_raw(item, path))


def get_path_raw(item: Any, path: str) -> Any:
    current = item
    for part in path.split("."):
        if current is None:
            return None
        if isinstance(current, dict):
            current = current.get(part)
            continue
        if isinstance(current, list):
            values = []
            for entry in current:
                if isinstance(entry, dict):
                    values.append(entry.get(part))
            current = values
            continue
        return None
    return current


def write_csv(path: Path, items: list[dict[str, Any]], columns: list[str]) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.DictWriter(fh, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        for item in items:
            writer.writerow(_flatten_row(item, columns))


def write_xlsx(path: Path, items: list[dict[str, Any]], columns: list[str]) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("XLSX export requires openpyxl. Use CSV or install openpyxl.") from exc

    wb = Workbook()
    ws = wb.active
    ws.title = "Assets"
    ws.append(columns)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")

    for item in items:
        row = _flatten_row(item, columns)
        ws.append([row.get(column, "") for column in columns])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for index, column in enumerate(columns, start=1):
        width = min(max(len(column) + 2, 12), 42)
        ws.column_dimensions[get_column_letter(index)].width = width
    wb.save(path)


def _flatten_row(item: dict[str, Any], columns: list[str]) -> dict[str, Any]:
    row: dict[str, Any] = {}
    for column in columns:
        value = item.get(column, "")
        if isinstance(value, (dict, list)):
            row[column] = str(value)
        elif value is None:
            row[column] = ""
        else:
            row[column] = value
    return row


def _stringify(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, dict):
        for key in ("name", "completename", "display_name", "version", "id"):
            if key in value and value[key] not in (None, ""):
                return str(value[key])
        return str(value)
    if isinstance(value, list):
        return ", ".join(_stringify(item) for item in value if _stringify(item))
    return str(value)
